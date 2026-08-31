"""Quick refresh: re-price the portfolio and rebuild the deliverable (~2 min).

Only the snapshot entries for the codes actually held in the workbook are
re-fetched; the annual/statement CSVs are left untouched. The rebuilt file is
recalculated by LibreOffice and verified in a temp dir, and only replaces the
deliverable once verification passes - a failed refresh never damages the file
the user already has.

refresh_full.py imports the shared pipeline from here.
"""

import os
import shutil
import subprocess
import sys
import tempfile
import time

from openpyxl import load_workbook

import builder
from fetcher import fetch_quick

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.normpath(os.path.join(HERE, "..", "Sahm_Portfolio_Analysis_v2.xlsx"))
DATA_DIR = os.path.join(HERE, "data")
VERIFY = os.path.join(HERE, "verify_workbook.py")

SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"
LO_TIMEOUT_S = 600

# Portfolio!B6:B25 - the editable holdings block.
PF_FIRST_ROW = 6
PF_LAST_ROW = 25

DEMO_CODES = ["2222", "1120", "2010", "7010"]

# Arabic console output must survive a redirected (non-UTF-8) stdout.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass


def portfolio_codes(path=OUT):
    """The 4-digit codes held in Portfolio!B6:B25 of `path`, else DEMO_CODES.

    Any unreadable workbook, missing sheet or empty holdings block degrades to
    the demo list rather than raising - a refresh must always have something to
    fetch.
    """
    if not os.path.exists(path):
        print("no deliverable yet at %s; using demo holdings" % path)
        return list(DEMO_CODES)
    try:
        wb = load_workbook(path, data_only=False)
    except Exception as exc:                    # noqa: BLE001 - any xlsx defect
        print("WARNING: could not read %s (%s); using demo holdings" % (path, exc))
        return list(DEMO_CODES)
    if "Portfolio" not in wb.sheetnames:
        print("WARNING: %s has no Portfolio sheet; using demo holdings" % path)
        wb.close()
        return list(DEMO_CODES)

    ws = wb["Portfolio"]
    codes = []
    for row in range(PF_FIRST_ROW, PF_LAST_ROW + 1):
        value = ws.cell(row=row, column=2).value
        if value is None:
            continue
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        code = str(value).strip()
        if not code:
            continue
        if len(code) != 4 or not code.isdigit():
            print("WARNING: ignoring non-code %r in Portfolio!B%d" % (code, row))
            continue
        if code not in codes:
            codes.append(code)
    wb.close()

    if not codes:
        print("Portfolio holdings are empty; using demo holdings")
        return list(DEMO_CODES)
    print("portfolio holds %d symbol(s): %s" % (len(codes), ", ".join(codes)))
    return codes


def _as_uri(path):
    """file:///C:/... form LibreOffice wants for -env:UserInstallation."""
    return "file:///" + os.path.abspath(path).replace("\\", "/").lstrip("/")


def lo_recalc(src, workdir):
    """Round-trip `src` through LibreOffice so formulas carry cached values.

    Returns the path of the recalculated copy (in its own dir - LibreOffice
    refuses to convert onto its own input) or None if the conversion failed.
    """
    if not os.path.exists(SOFFICE):
        print("FAIL  LibreOffice not found at %s" % SOFFICE)
        return None

    outdir = os.path.join(workdir, "lo_out")
    profile = os.path.join(workdir, "lo_profile")
    os.makedirs(outdir, exist_ok=True)
    os.makedirs(profile, exist_ok=True)

    cmd = [
        SOFFICE, "--headless", "--norestore", "--invisible",
        "-env:UserInstallation=%s" % _as_uri(profile),
        "--convert-to", "xlsx:Calc MS Excel 2007 XML",
        "--outdir", outdir, src,
    ]
    print("recalculating via LibreOffice ...")
    try:
        proc = subprocess.run(cmd, capture_output=True, text=True,
                              timeout=LO_TIMEOUT_S)
    except subprocess.TimeoutExpired:
        print("FAIL  LibreOffice recalc timed out after %ds" % LO_TIMEOUT_S)
        return None

    out = (proc.stdout or "") + (proc.stderr or "")
    if out.strip():
        print(out.strip())

    recalced = os.path.join(outdir, os.path.basename(src))
    if proc.returncode != 0 or not os.path.exists(recalced):
        print("FAIL  LibreOffice recalc failed (exit %d)" % proc.returncode)
        return None
    return recalced


def verify(path):
    """Run verify_workbook.py on `path`; returns (ok, combined output)."""
    proc = subprocess.run(
        [sys.executable, VERIFY, path, "--data-dir", DATA_DIR],
        capture_output=True, text=True, encoding="utf-8", errors="replace")
    return proc.returncode == 0, (proc.stdout or "") + (proc.stderr or "")


def promote(recalced):
    """Copy the verified build over the deliverable. True when it landed."""
    try:
        shutil.copy2(recalced, OUT)
    except PermissionError:
        print("FAIL  %s is locked (open in Excel?) - close it and re-run" % OUT)
        return False
    print("wrote %s" % OUT)
    return True


def pipeline(fetch):
    """fetch -> build -> recalc -> verify -> promote.

    `fetch` is a zero-argument callable doing the data refresh. Everything is
    staged in a temp dir; the deliverable is only overwritten after verify
    passes. Returns elapsed seconds on success, None on failure.
    """
    start = time.time()
    fetch()

    workdir = tempfile.mkdtemp(prefix="sahm_v2_")
    try:
        stage = os.path.join(workdir, os.path.basename(OUT))
        preserve = OUT if os.path.exists(OUT) else None
        written = builder.build(stage, data_dir=DATA_DIR, preserve_from=preserve)

        recalced = lo_recalc(written, workdir)
        if recalced is None:
            print("previous deliverable kept unchanged.")
            return None

        ok, output = verify(recalced)
        print(output.rstrip())
        if not ok:
            print("FAIL  verification failed - previous deliverable kept unchanged.")
            return None

        if not promote(recalced):
            return None
    finally:
        shutil.rmtree(workdir, ignore_errors=True)

    return time.time() - start


def main():
    codes = portfolio_codes()
    elapsed = pipeline(lambda: fetch_quick(codes, data_dir=DATA_DIR))
    if elapsed is None:
        print("✖ فشل التحديث — لم يتم تعديل الملف")
        return 1
    print("✔ تم تحديث المحفظة (%d رمز، %.0f ثانية)" % (len(codes), elapsed))
    return 0


if __name__ == "__main__":
    sys.exit(main())
