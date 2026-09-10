"""Generate raw Google Sheets batchUpdate requests for the live Waraqah book.

This script never authenticates to Google and never writes remotely.  It turns
the verified local workbook into three auditable request phases consumed by the
connected Google Drive tool: structure, values, and format/validation.
"""

import argparse
from datetime import date, datetime
import json

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


SHEET_IDS = {
    "Portfolio": 580439390,
    "Orders": 2000000101,
    "Stock Lookup": 690017234,
    "Performance": 2000000102,
    "Activity": 2000000103,
    "Sharia": 2000000104,
    "Risk & Horizons": 1697126793,
    "DB": 640167122,
    "Statements": 145447140,
    "Symbols": 1932490456,
    "Checks": 2000000105,
    "Guide": 1251951118,
}

ORDER = list(SHEET_IDS)
NEW_SHEETS = {
    "Orders": (600, 57, 5),
    "Performance": (600, 9, 5),
    "Activity": (600, 28, 5),
    "Sharia": (600, 23, 5),
    "Checks": (600, 6, 4),
}

BLUE = {"red": 0.12156863, "green": 0.21960784, "blue": 0.39215687}
PALE_BLUE = {"red": 0.8509804, "green": 0.8862745, "blue": 0.9529412}
GRAY = {"red": 0.8509804, "green": 0.8509804, "blue": 0.8509804}
PALE_YELLOW = {"red": 1.0, "green": 0.9490196, "blue": 0.8}
WHITE = {"red": 1.0, "green": 1.0, "blue": 1.0}
GREEN = {"red": 0.0, "green": 0.38039216, "blue": 0.0}
RED = {"red": 0.7529412, "green": 0.0, "blue": 0.0}
AMBER = {"red": 0.75, "green": 0.5, "blue": 0.0}
PALE_GREEN = {"red": 0.85, "green": 0.94, "blue": 0.83}
PALE_RED = {"red": 0.96, "green": 0.80, "blue": 0.80}
PALE_AMBER = {"red": 1.0, "green": 0.90, "blue": 0.65}
BORDER = {"style": "SOLID", "color": {"red": 0.749, "green": 0.749, "blue": 0.749}}


def grid(sheet_id, r0, r1, c0, c1):
    return {"sheetId": sheet_id, "startRowIndex": r0, "endRowIndex": r1,
            "startColumnIndex": c0, "endColumnIndex": c1}


def repeat(sheet_id, r0, r1, c0, c1, cell_format):
    return {"repeatCell": {"range": grid(sheet_id, r0, r1, c0, c1),
                           "cell": {"userEnteredFormat": cell_format},
                           "fields": "userEnteredFormat"}}


def dimension(sheet_id, dimension_name, start, end, pixels):
    return {"updateDimensionProperties": {
        "range": {"sheetId": sheet_id, "dimension": dimension_name,
                  "startIndex": start, "endIndex": end},
        "properties": {"pixelSize": pixels}, "fields": "pixelSize"}}


def entered_value(value, data_type):
    if value is None:
        return {}
    if data_type == "f" or (isinstance(value, str) and value.startswith("=")):
        return {"userEnteredValue": {"formulaValue": value}}
    if isinstance(value, bool):
        return {"userEnteredValue": {"boolValue": value}}
    if isinstance(value, (int, float)):
        return {"userEnteredValue": {"numberValue": value}}
    if isinstance(value, (date, datetime)):
        value = value.isoformat()
    return {"userEnteredValue": {"stringValue": str(value)}}


def value_requests(workbook, chunk_rows=150, only_sheet=None,
                   requested_first=None, requested_last=None,
                   requested_first_col=None, requested_last_col=None):
    requests = []
    # The new guide is shorter than the imported one. Clear old text after the
    # structure phase has removed its merges, then write the current guide.
    names = [only_sheet] if only_sheet else ORDER
    if (only_sheet in (None, "Guide") and
            (requested_first is None or requested_first <= 1)):
        requests.append({"updateCells": {
            "range": grid(SHEET_IDS["Guide"], 0, 100, 0, 6),
            "fields": "userEnteredValue"}})
    for name in names:
        ws = workbook[name]
        max_row, max_col = ws.max_row, ws.max_column
        first_bound = max(1, requested_first or 1)
        last_bound = min(max_row, requested_last or max_row)
        first_col = max(1, requested_first_col or 1)
        last_col = min(max_col, requested_last_col or max_col)
        for first in range(first_bound, last_bound + 1, chunk_rows):
            last = min(last_bound, first + chunk_rows - 1)
            rows = []
            for row in ws.iter_rows(min_row=first, max_row=last,
                                    min_col=first_col, max_col=last_col):
                rows.append({"values": [entered_value(cell.value, cell.data_type)
                                        for cell in row]})
            requests.append({"updateCells": {
                "start": {"sheetId": SHEET_IDS[name], "rowIndex": first - 1,
                          "columnIndex": first_col - 1},
                "rows": rows, "fields": "userEnteredValue"}})
    return requests


def structure_requests():
    requests = [{"updateSpreadsheetProperties": {
        "properties": {"timeZone": "Asia/Riyadh"}, "fields": "timeZone"}}]
    # Insert in final-order positions; each insertion shifts the legacy tabs.
    for name in ("Orders", "Performance", "Activity", "Sharia", "Checks"):
        rows, cols, frozen = NEW_SHEETS[name]
        requests.append({"addSheet": {"properties": {
            "sheetId": SHEET_IDS[name], "title": name, "index": ORDER.index(name),
            "gridProperties": {"rowCount": rows, "columnCount": cols,
                               "frozenRowCount": frozen, "hideGridlines": True}}}})
    requests.extend([
        {"updateSheetProperties": {
            "properties": {"sheetId": SHEET_IDS["Portfolio"],
                           "gridProperties": {"columnCount": 36, "frozenRowCount": 5,
                                              "hideGridlines": True}},
            "fields": "gridProperties.columnCount,gridProperties.frozenRowCount,gridProperties.hideGridlines"}},
        {"updateSheetProperties": {
            "properties": {"sheetId": SHEET_IDS["Risk & Horizons"],
                           "gridProperties": {"columnCount": 47, "frozenRowCount": 4,
                                              "hideGridlines": True}},
            "fields": "gridProperties.columnCount,gridProperties.frozenRowCount,gridProperties.hideGridlines"}},
        {"unmergeCells": {"range": grid(SHEET_IDS["Guide"], 0, 100, 0, 6)}},
    ])
    return requests


def format_requests(workbook, include_guide_merges=True):
    requests = []
    table_specs = {
        "Orders": (5, 505, 57), "Activity": (5, 505, 28),
        "Sharia": (5, 505, 23), "Performance": (30, 505, 9),
    }
    title_fmt = {"textFormat": {"fontFamily": "Arial", "fontSize": 14,
                                 "bold": True, "foregroundColor": BLUE},
                 "verticalAlignment": "MIDDLE"}
    note_fmt = {"textFormat": {"fontFamily": "Arial", "fontSize": 9,
                                "italic": True, "foregroundColor": {"red": .35, "green": .35, "blue": .35}},
                "wrapStrategy": "WRAP", "verticalAlignment": "MIDDLE"}
    title_left_fmt = {**title_fmt, "horizontalAlignment": "LEFT",
                      "wrapStrategy": "OVERFLOW_CELL"}
    note_left_fmt = {**note_fmt, "horizontalAlignment": "LEFT",
                     "wrapStrategy": "OVERFLOW_CELL"}
    header_fmt = {"backgroundColor": BLUE,
                  "textFormat": {"fontFamily": "Arial", "fontSize": 10,
                                 "bold": True, "foregroundColor": WHITE},
                  "horizontalAlignment": "CENTER", "verticalAlignment": "MIDDLE",
                  "wrapStrategy": "WRAP",
                  "borders": {side: BORDER for side in ("top", "bottom", "left", "right")}}
    body_fmt = {"textFormat": {"fontFamily": "Arial", "fontSize": 10},
                "verticalAlignment": "MIDDLE", "wrapStrategy": "WRAP",
                "borders": {side: BORDER for side in ("top", "bottom", "left", "right")}}
    blank_fmt = {"backgroundColor": WHITE,
                 "textFormat": {"fontFamily": "Arial", "fontSize": 10},
                 "verticalAlignment": "MIDDLE", "wrapStrategy": "OVERFLOW_CELL"}
    input_fmt = {
        **body_fmt, "backgroundColor": PALE_YELLOW,
        "textFormat": {"fontFamily": "Arial", "fontSize": 10, "bold": True,
                       "foregroundColor": {"red": 0.0, "green": 0.0, "blue": 1.0}},
    }

    for name in NEW_SHEETS:
        sid = SHEET_IDS[name]
        requests.append(repeat(sid, 1, 2, 0, NEW_SHEETS[name][1], title_fmt))
        requests.append(repeat(sid, 2, 3, 0, NEW_SHEETS[name][1], note_fmt))
        requests.append(repeat(sid, 1, 2, 0, 1, title_left_fmt))
        requests.append(repeat(sid, 2, 3, 0, 1, note_left_fmt))
        requests.append(dimension(sid, "ROWS", 1, 2, 28))
        requests.append(dimension(sid, "ROWS", 2, 3, 52))

    # Imported Arabic titles otherwise inherit RIGHT alignment in the leftmost
    # cell and overflow off-screen.  Anchor them left while keeping their font.
    for name, r0, c0 in [
        ("Portfolio", 1, 1), ("Stock Lookup", 0, 1),
        ("Risk & Horizons", 1, 0), ("DB", 1, 0),
        ("Statements", 0, 0), ("Symbols", 1, 0),
    ]:
        requests.append(repeat(SHEET_IDS[name], r0, r0 + 1, c0, c0 + 1,
                               title_left_fmt))
    requests.append(repeat(SHEET_IDS["Portfolio"], 2, 3, 1, 2,
                           note_left_fmt))

    for name, (header_row, last_row, cols) in table_specs.items():
        sid = SHEET_IDS[name]
        requests.append(repeat(sid, header_row - 1, header_row, 0, cols, header_fmt))
        requests.append(repeat(sid, header_row, last_row, 0, cols, body_fmt))
        requests.append({"setBasicFilter": {"filter": {"range": grid(
            sid, header_row - 1, last_row, 0, cols)}}})
        requests.append(dimension(sid, "ROWS", header_row - 1, header_row, 40))

    # Checks has a settings block, a checks table and a review log.
    sid = SHEET_IDS["Checks"]
    requests.extend([
        repeat(sid, 3, 4, 0, 3, header_fmt), repeat(sid, 4, 20, 0, 3, body_fmt),
        repeat(sid, 22, 23, 0, 4, header_fmt), repeat(sid, 23, 41, 0, 4, body_fmt),
        repeat(sid, 43, 44, 0, 6, header_fmt), repeat(sid, 44, 45, 0, 6, body_fmt),
        repeat(sid, 4, 20, 1, 2, input_fmt),
    ])

    # Existing sheets keep their imported style; format only newly extended columns.
    requests.extend([
        repeat(SHEET_IDS["Portfolio"], 4, 5, 14, 36,
               {**header_fmt, "backgroundColor": GRAY,
                "textFormat": {"fontFamily": "Arial", "fontSize": 10,
                               "bold": True, "foregroundColor": BLUE}}),
        repeat(SHEET_IDS["Portfolio"], 5, 26, 14, 36, body_fmt),
        repeat(SHEET_IDS["Risk & Horizons"], 3, 4, 19, 47, header_fmt),
        repeat(SHEET_IDS["Risk & Horizons"], 4, 206, 19, 47,
               {**body_fmt, "textFormat": {"fontFamily": "Arial", "fontSize": 10,
                                           "foregroundColor": GREEN}}),
        {"setBasicFilter": {"filter": {"range": grid(
            SHEET_IDS["Portfolio"], 4, 25, 0, 36)}}},
        {"setBasicFilter": {"filter": {"range": grid(
            SHEET_IDS["Risk & Horizons"], 3, 206, 0, 47)}}},
    ])

    # Input areas are visible but calculated columns stay neutral.
    requests.extend([
        repeat(SHEET_IDS["Performance"], 4, 29, 3, 9, blank_fmt),
        repeat(SHEET_IDS["Performance"], 4, 5, 0, 3, header_fmt),
        repeat(SHEET_IDS["Performance"], 5, 27, 0, 3, body_fmt),
        repeat(SHEET_IDS["Activity"], 5, 505, 0, 17, input_fmt),
        repeat(SHEET_IDS["Activity"], 5, 505, 24, 26, input_fmt),
        repeat(SHEET_IDS["Activity"], 5, 505, 27, 28, input_fmt),
        repeat(SHEET_IDS["Sharia"], 5, 505, 4, 21, input_fmt),
        repeat(SHEET_IDS["Orders"], 5, 505, 11, 12, input_fmt),
        repeat(SHEET_IDS["Orders"], 5, 505, 13, 20, input_fmt),
        repeat(SHEET_IDS["Orders"], 5, 505, 28, 30, input_fmt),
        repeat(SHEET_IDS["Orders"], 5, 505, 43, 54, input_fmt),
        repeat(SHEET_IDS["Performance"], 30, 505, 0, 3, input_fmt),
        repeat(SHEET_IDS["Performance"], 30, 505, 5, 9, input_fmt),
    ])

    # Number formats for key accounting columns.
    for name, r0, r1, c0, c1, pattern in [
        ("Activity", 5, 505, 9, 11, "#,##0.00"),
        ("Activity", 5, 505, 12, 16, "#,##0.00"),
        ("Activity", 5, 505, 17, 24, "#,##0.00"),
        ("Orders", 5, 505, 16, 17, "#,##0.00"),
        ("Orders", 5, 505, 17, 19, "#,##0.00"),
        ("Orders", 5, 505, 30, 33, "0.0%"),
        ("Orders", 5, 505, 33, 36, "#,##0.00"),
        ("Orders", 5, 505, 47, 48, "#,##0.00"),
        ("Orders", 5, 505, 48, 49, "#,##0.00"),
        ("Performance", 5, 17, 1, 2, "#,##0.00"),
        ("Performance", 17, 18, 1, 2, "0.0%"),
        ("Performance", 20, 22, 1, 2, "0.0%"),
        ("Performance", 22, 24, 1, 2, "#,##0.00"),
        ("Performance", 30, 505, 1, 3, "#,##0.00"),
        ("Performance", 30, 505, 3, 4, "0.0%"),
        ("Performance", 30, 505, 4, 5, "0.00"),
        ("Performance", 30, 505, 6, 7, "0.0%"),
    ]:
        requests.append({"repeatCell": {
            "range": grid(SHEET_IDS[name], r0, r1, c0, c1),
            "cell": {"userEnteredFormat": {
                "numberFormat": {"type": "NUMBER", "pattern": pattern}}},
            "fields": "userEnteredFormat.numberFormat"}})

    # Controlled vocabularies.
    validations = [
        ("Activity", 5, 505, 7, 8, ["Tadawul", "NYSE", "NASDAQ", "Other"]),
        ("Activity", 5, 505, 8, 9, ["Opening balance", "Buy", "Sell", "Dividend", "Fee", "Tax", "Withholding", "Deposit", "Withdrawal", "Transfer in", "Transfer out", "Split", "Corporate action", "FX buy", "FX sell"]),
        ("Activity", 5, 505, 11, 12, ["SAR", "USD"]),
        ("Sharia", 5, 505, 4, 5, ["Compliant", "Non-compliant", "Uncertain", "Review overdue"]),
        ("Orders", 5, 505, 5, 6, ["6 months", "2 years", "5 years"]),
        ("Orders", 5, 505, 13, 14, ["Buy", "Add", "Hold", "Trim", "Exit", "Wait"]),
        ("Orders", 5, 505, 14, 15, ["Limit", "Stop", "Stop limit", "Market", "No order"]),
        ("Orders", 5, 505, 42, 43, ["Low", "Medium", "High"]),
        ("Orders", 5, 505, 43, 44, ["Proposed", "Approved", "Submitted", "Part-filled", "Filled", "Cancelled", "Expired", "Superseded"]),
        ("Orders", 5, 505, 44, 45, ["Approve", "Reject", "Modify", "Defer"]),
    ]
    for name, r0, r1, c0, c1, values in validations:
        requests.append({"setDataValidation": {
            "range": grid(SHEET_IDS[name], r0, r1, c0, c1),
            "rule": {"condition": {"type": "ONE_OF_LIST",
                                    "values": [{"userEnteredValue": value} for value in values]},
                     "strict": True, "showCustomUi": True}}})

    # Status highlighting.
    rules = [
        ("Sharia", 5, 505, 4, 5, "Compliant", PALE_GREEN, GREEN),
        ("Sharia", 5, 505, 4, 5, "Non-compliant", PALE_RED, RED),
        ("Sharia", 5, 505, 4, 5, "Uncertain", PALE_AMBER, AMBER),
        ("Sharia", 5, 505, 4, 5, "Review overdue", PALE_RED, RED),
        ("Checks", 23, 41, 1, 2, "PASS", PALE_GREEN, GREEN),
        ("Checks", 23, 41, 1, 2, "FAIL", PALE_RED, RED),
        ("Checks", 23, 41, 1, 2, "WARN", PALE_AMBER, AMBER),
    ]
    for name, r0, r1, c0, c1, value, bg, fg in rules:
        requests.append({"addConditionalFormatRule": {"index": 0, "rule": {
            "ranges": [grid(SHEET_IDS[name], r0, r1, c0, c1)],
            "booleanRule": {"condition": {"type": "TEXT_EQ",
                                           "values": [{"userEnteredValue": value}]},
                            "format": {"backgroundColor": bg,
                                       "textFormat": {"foregroundColor": fg, "bold": True}}}}}})

    # Widths tuned for scannability.  Unspecified new columns default to 120px.
    for name, (_, cols, _) in NEW_SHEETS.items():
        for col in range(cols):
            requests.append(dimension(SHEET_IDS[name], "COLUMNS", col, col + 1, 120))
    wide = {
        "Orders": {
            0: 220, 8: 220, 15: 270, 20: 340, 21: 390, 22: 330, 23: 330,
            24: 300, 25: 280, 26: 320, 27: 320, 28: 250, 29: 270,
            36: 340, 37: 340, 39: 300, 41: 380, 46: 190, 49: 220,
            50: 190, 53: 300, 55: 220, 56: 190,
        },
        "Activity": {0: 190, 3: 140, 4: 140, 5: 130, 8: 140, 16: 170, 24: 280, 27: 300},
        "Sharia": {
            2: 220, 5: 190, 6: 190, 7: 170, 8: 300, 9: 170,
            12: 260, 13: 260, 14: 240, 18: 220, 20: 300, 22: 210,
        },
        "Performance": {0: 190, 2: 270, 5: 170, 7: 190, 8: 300},
        "Checks": {0: 220, 1: 180, 2: 300, 3: 340, 4: 340, 5: 340},
    }
    for name, mapping in wide.items():
        for col, pixels in mapping.items():
            requests.append(dimension(SHEET_IDS[name], "COLUMNS", col, col + 1, pixels))
    for col in range(14, 36):
        requests.append(dimension(SHEET_IDS["Portfolio"], "COLUMNS", col, col + 1,
                                  150 if col not in (26, 27, 28, 29, 30, 31, 32, 33)
                                  else (300 if col in (26, 27) else 190)))
    requests.append(dimension(SHEET_IDS["Portfolio"], "COLUMNS", 13, 14, 200))
    requests.append(dimension(SHEET_IDS["Portfolio"], "COLUMNS", 26, 27, 220))
    for col in range(19, 47):
        requests.append(dimension(SHEET_IDS["Risk & Horizons"], "COLUMNS", col, col + 1,
                                  150 if col not in (24, 30, 34, 35, 36, 37, 38,
                                                    39, 40, 41, 42, 43, 44, 45, 46)
                                  else (300 if col in (24, 38, 39, 43, 44, 46) else 190)))

    # Rebuild the Guide's intentional merge topology after values are written.
    guide = workbook["Guide"]
    if include_guide_merges:
        for merged in guide.merged_cells.ranges:
            min_col, min_row, max_col, max_row = range_boundaries(str(merged))
            requests.append({"mergeCells": {
                "range": grid(SHEET_IDS["Guide"], min_row - 1, max_row,
                              min_col - 1, max_col), "mergeType": "MERGE_ALL"}})
    requests.extend([
        repeat(SHEET_IDS["Guide"], 0, guide.max_row, 1, 6,
               {"textFormat": {"fontFamily": "Arial", "fontSize": 10},
                "horizontalAlignment": "RIGHT", "verticalAlignment": "TOP",
                "wrapStrategy": "WRAP"}),
        dimension(SHEET_IDS["Guide"], "COLUMNS", 1, 6, 210),
    ])
    for merged in guide.merged_cells.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(merged))
        anchor = guide.cell(min_row, min_col)
        if anchor.font.bold:
            fmt = {"backgroundColor": PALE_BLUE,
                   "textFormat": {"fontFamily": "Arial", "fontSize": 12,
                                  "bold": True, "foregroundColor": BLUE},
                   "horizontalAlignment": "RIGHT",
                   "verticalAlignment": "MIDDLE",
                   "wrapStrategy": "WRAP"}
            height = 30
        else:
            fmt = {"backgroundColor": WHITE,
                   "textFormat": {"fontFamily": "Arial", "fontSize": 10,
                                  "bold": False},
                   "horizontalAlignment": "RIGHT",
                   "verticalAlignment": "TOP",
                   "wrapStrategy": "WRAP"}
            height = 46
        requests.append(repeat(SHEET_IDS["Guide"], min_row - 1, max_row,
                               min_col - 1, max_col, fmt))
        requests.append(dimension(SHEET_IDS["Guide"], "ROWS", min_row - 1,
                                  max_row, height))
    return requests


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("--phase", choices=("structure", "values", "format"), required=True)
    parser.add_argument("--sheet", choices=ORDER)
    parser.add_argument("--first-row", type=int)
    parser.add_argument("--last-row", type=int)
    parser.add_argument("--first-column", type=int)
    parser.add_argument("--last-column", type=int)
    parser.add_argument("--omit-guide-merges", action="store_true",
                        help="format an already-merged live Guide without re-merging")
    args = parser.parse_args()
    wb = load_workbook(args.workbook, data_only=False)
    if args.phase == "structure":
        requests = structure_requests()
    elif args.phase == "values":
        requests = value_requests(wb, only_sheet=args.sheet,
                                  requested_first=args.first_row,
                                  requested_last=args.last_row,
                                  requested_first_col=args.first_column,
                                  requested_last_col=args.last_column)
    else:
        requests = format_requests(wb, include_guide_merges=not args.omit_guide_merges)
    print(json.dumps(requests, ensure_ascii=False, separators=(",", ":")))
    wb.close()


if __name__ == "__main__":
    main()
